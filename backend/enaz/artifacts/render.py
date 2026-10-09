"""Deterministic renderers: spec → .pptx / .docx / .xlsx bytes.

Numbers come only from the spec (which is built from tool/evidence output), never
invented at render time. Citations ride along: speaker notes in decks, a Sources
section in docs, a Sources tab in sheets.
"""

from __future__ import annotations

import io

from .specs import DeckSpec, DocSpec, SheetSpec

BRAND = "4F46E5"  # indigo accent


def render_deck(spec: DeckSpec) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_title(slide, text, top=0.5, size=34):
        box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(1.2))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = text
        p = tf.paragraphs[0]
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = RGBColor.from_string(BRAND)
        return box

    # Title slide
    title_slide = prs.slides.add_slide(blank)
    add_title(title_slide, spec.title, top=2.6, size=40)
    if spec.subtitle:
        sb = title_slide.shapes.add_textbox(Inches(0.7), Inches(3.9), Inches(12), Inches(1))
        sb.text_frame.text = spec.subtitle
        sb.text_frame.paragraphs[0].font.size = Pt(18)

    for slide_spec in spec.slides:
        slide = prs.slides.add_slide(blank)
        add_title(slide, slide_spec.title)
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(12), Inches(5.2))
        tf = body.text_frame
        tf.word_wrap = True

        if slide_spec.layout == "chart" and slide_spec.chart_values:
            from pptx.chart.data import CategoryChartData
            from pptx.enum.chart import XL_CHART_TYPE

            data = CategoryChartData()
            data.categories = slide_spec.chart_labels or [str(i + 1) for i in range(len(slide_spec.chart_values))]
            data.add_series("Series 1", slide_spec.chart_values)
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(11), Inches(4.8), data
            )
        elif slide_spec.layout == "table" and slide_spec.rows:
            cols = len(slide_spec.columns) or len(slide_spec.rows[0])
            rows = len(slide_spec.rows) + (1 if slide_spec.columns else 0)
            table = slide.shapes.add_table(rows, cols, Inches(0.7), Inches(1.9), Inches(12), Inches(0.4 * rows)).table
            r0 = 0
            if slide_spec.columns:
                for c, head in enumerate(slide_spec.columns):
                    table.cell(0, c).text = str(head)
                r0 = 1
            for ri, row in enumerate(slide_spec.rows):
                for ci, val in enumerate(row[:cols]):
                    table.cell(ri + r0, ci).text = str(val)
        else:
            first = True
            for bullet in slide_spec.bullets:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                p.text = f"•  {bullet}"
                p.font.size = Pt(18)
                first = False

        if slide_spec.notes:
            slide.notes_slide.notes_text_frame.text = slide_spec.notes

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def render_doc(spec: DocSpec) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    h = document.add_heading(spec.title, level=0)
    if h.runs:
        h.runs[0].font.color.rgb = RGBColor.from_string(BRAND)
    if spec.subtitle:
        document.add_paragraph(spec.subtitle).italic = True

    for section in spec.sections:
        if section.heading:
            document.add_heading(section.heading, level=1)
        for para in section.paragraphs:
            document.add_paragraph(para)
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
        if section.rows:
            cols = len(section.columns) or len(section.rows[0])
            table = document.add_table(rows=0, cols=cols)
            table.style = "Light Grid Accent 1"
            if section.columns:
                cells = table.add_row().cells
                for c, head in enumerate(section.columns):
                    cells[c].text = str(head)
            for row in section.rows:
                cells = table.add_row().cells
                for c, val in enumerate(row[:cols]):
                    cells[c].text = str(val)

    if spec.sources:
        document.add_heading("Sources", level=1)
        for i, src in enumerate(spec.sources, start=1):
            p = document.add_paragraph(f"{i}. {src}")
            p.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def render_sheet(spec: SheetSpec) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor=BRAND)
    header_font = Font(bold=True, color="FFFFFF")

    for tab in spec.tabs or []:
        ws = wb.create_sheet(title=(tab.name or "Sheet")[:31])
        headers = list(tab.columns)
        if tab.formula_header:
            headers = headers + [tab.formula_header]
        for c, head in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=head)
            cell.fill = header_fill
            cell.font = header_font
        numeric_cols: set[int] = set()
        for r, row in enumerate(tab.rows, start=2):
            for c, val in enumerate(row, start=1):
                num = _as_number(val)
                ws.cell(row=r, column=c, value=num if num is not None else val)
                if num is not None:
                    numeric_cols.add(c)
            if tab.formula_template:
                formula = tab.formula_template.replace("{r}", str(r))
                ws.cell(row=r, column=len(tab.columns) + 1, value=formula if formula.startswith("=") else f"={formula}")
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 22
        # A simple bar chart over the first numeric column, when present.
        if tab.rows and numeric_cols:
            col = min(numeric_cols)
            chart = BarChart()
            chart.title = tab.name
            data = Reference(ws, min_col=col, min_row=1, max_row=len(tab.rows) + 1)
            cats = Reference(ws, min_col=1, min_row=2, max_row=len(tab.rows) + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, f"{get_column_letter(len(headers) + 3)}2")

    if spec.sources:
        ws = wb.create_sheet(title="Sources")
        ws.cell(row=1, column=1, value="Sources").font = Font(bold=True)
        for i, src in enumerate(spec.sources, start=2):
            ws.cell(row=i, column=1, value=src)
        ws.column_dimensions["A"].width = 80
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _as_number(value) -> float | int | None:
    if isinstance(value, (int, float)):
        return value
    try:
        s = str(value).replace(",", "").replace("$", "").replace("%", "").strip()
        f = float(s)
        return int(f) if f.is_integer() else f
    except (ValueError, AttributeError):
        return None


def render(kind: str, spec) -> bytes:
    if kind == "slides":
        return render_deck(spec)
    if kind == "doc":
        return render_doc(spec)
    if kind == "sheet":
        return render_sheet(spec)
    raise ValueError(f"Unknown artifact kind: {kind}")
