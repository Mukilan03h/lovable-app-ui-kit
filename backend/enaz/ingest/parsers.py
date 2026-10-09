"""Turn files into structured text: a list of sections, each with a heading path.

Keeping headings (and slide / sheet / page boundaries) lets the chunker split on
structure and lets citations point at a precise section.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath


@dataclass
class Section:
    heading: str
    text: str
    level: int = 1


@dataclass
class ParsedDocument:
    title: str
    sections: list[Section] = field(default_factory=list)
    doc_type: str = "doc"
    metadata: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(f"{s.heading}\n{s.text}" if s.heading else s.text for s in self.sections)


class UnsupportedFile(ValueError):
    pass


EXTENSION_TYPES = {
    ".pdf": "pdf",
    ".docx": "doc",
    ".pptx": "slides",
    ".xlsx": "sheet",
    ".csv": "sheet",
    ".md": "page",
    ".markdown": "page",
    ".txt": "doc",
    ".html": "page",
    ".htm": "page",
    ".json": "doc",
}


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    ext = PurePath(filename).suffix.lower()
    title = PurePath(filename).stem.replace("_", " ").replace("-", " ").strip() or filename
    if ext == ".pdf":
        return _parse_pdf(title, data)
    if ext == ".docx":
        return _parse_docx(title, data)
    if ext == ".pptx":
        return _parse_pptx(title, data)
    if ext == ".xlsx":
        return _parse_xlsx(title, data)
    if ext == ".csv":
        return _parse_csv(title, data.decode("utf-8", errors="replace"))
    if ext in (".md", ".markdown"):
        return parse_markdown(title, data.decode("utf-8", errors="replace"))
    if ext in (".html", ".htm"):
        return parse_html(data.decode("utf-8", errors="replace"), fallback_title=title)
    if ext in (".txt", ".json", ""):
        return parse_text(title, data.decode("utf-8", errors="replace"))
    raise UnsupportedFile(f"Unsupported file type: {ext or 'unknown'}")


def parse_text(title: str, text: str, doc_type: str = "doc") -> ParsedDocument:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return ParsedDocument(title=title, sections=[Section("", "\n\n".join(paragraphs))], doc_type=doc_type)


def parse_markdown(title: str, text: str) -> ParsedDocument:
    sections: list[Section] = []
    stack: list[str] = []
    buf: list[str] = []
    current_level = 1
    doc_title = title

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(" › ".join(stack), body, current_level))
        buf.clear()

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            level = len(m.group(1))
            heading = m.group(2).strip()
            if level == 1 and not sections and not stack:
                doc_title = heading
            stack[:] = stack[: level - 1] + [heading]
            current_level = level
        else:
            buf.append(line)
    flush()
    if not sections:
        sections.append(Section("", text.strip()))
    return ParsedDocument(title=doc_title, sections=sections, doc_type="page")


def parse_html(html: str, fallback_title: str = "Web page", url: str | None = None) -> ParsedDocument:
    from bs4 import BeautifulSoup

    # Prefer readability's main-content extraction (drops chrome/boilerplate);
    # fall back to structural parsing of the raw HTML.
    try:
        from readability import Document as ReadabilityDocument

        rd = ReadabilityDocument(html)
        main_html = rd.summary(html_partial=True)
        title = (rd.short_title() or fallback_title).strip()
        soup = BeautifulSoup(main_html, "html.parser")
        sections: list[Section] = []
        heading = ""
        buf: list[str] = []
        for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote", "td"]):
            text = el.get_text(" ", strip=True)
            if not text:
                continue
            if el.name in ("h1", "h2", "h3", "h4"):
                if buf:
                    sections.append(Section(heading, "\n".join(buf)))
                    buf = []
                heading = text
            else:
                buf.append(text)
        if buf:
            sections.append(Section(heading, "\n".join(buf)))
        if sections and sum(len(s.text) for s in sections) > 200:
            return ParsedDocument(title=title, sections=sections, doc_type="page",
                                  metadata={"url": url} if url else {})
    except Exception:
        pass

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg", "form"]):
        tag.decompose()
    title = (soup.title.string.strip() if soup.title and soup.title.string else "") or fallback_title
    root = soup.find("main") or soup.find("article") or soup.body or soup
    sections: list[Section] = []
    heading = ""
    buf: list[str] = []
    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "td", "blockquote"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            if buf:
                sections.append(Section(heading, "\n".join(buf)))
                buf = []
            heading = text
        else:
            buf.append(text)
    if buf:
        sections.append(Section(heading, "\n".join(buf)))
    if not sections:
        text = root.get_text(" ", strip=True)
        sections.append(Section("", text))
    meta = {"url": url} if url else {}
    return ParsedDocument(title=title, sections=sections, doc_type="page", metadata=meta)


def _parse_pdf(title: str, data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(f"Page {i}", text))
    meta_title = (reader.metadata.title if reader.metadata and reader.metadata.title else "") or title
    return ParsedDocument(title=meta_title, sections=sections, doc_type="pdf", metadata={"pages": len(reader.pages)})


def _parse_docx(title: str, data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    sections: list[Section] = []
    stack: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        if buf:
            sections.append(Section(" › ".join(stack), "\n".join(buf)))
            buf.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            stack[:] = stack[: level - 1] + [text]
        else:
            buf.append(text)
    flush()
    for t_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            sections.append(Section(f"Table {t_index}", "\n".join(rows)))
    core_title = document.core_properties.title or title
    return ParsedDocument(title=core_title, sections=sections, doc_type="doc")


def _parse_pptx(title: str, data: bytes) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape == getattr(slide.shapes, "title", None) and not slide_title:
                    slide_title = text
                else:
                    lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        if slide_title or lines:
            sections.append(Section(f"Slide {i}: {slide_title}".strip(": "), "\n".join(lines)))
    return ParsedDocument(title=title, sections=sections, doc_type="slides", metadata={"slides": len(prs.slides)})


def _parse_xlsx(title: str, data: bytes) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sections = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(" | ".join(cells))
            if len(rows) >= 2000:
                break
        if rows:
            sections.append(Section(f"Sheet {ws.title}", "\n".join(rows)))
    return ParsedDocument(title=title, sections=sections, doc_type="sheet")


def _parse_csv(title: str, text: str) -> ParsedDocument:
    rows = [" | ".join(r) for r in csv.reader(io.StringIO(text)) if any(r)]
    return ParsedDocument(title=title, sections=[Section("", "\n".join(rows[:5000]))], doc_type="sheet")
